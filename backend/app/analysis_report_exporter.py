"""Excel and PDF exports for ActuTrace analysis reports."""
from pathlib import Path
from typing import Any
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle


def export_excel(summary: dict[str, Any], findings: list[dict], impacts: list[dict], path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["ActuTrace Analysis Summary", ""])
    ws.append(["Run ID", summary.get("run_metadata", {}).get("run_id", "")])
    ws.append(["Generated at (UTC)", summary.get("run_metadata", {}).get("generated_at_utc", "")])
    ws.append(["Validation findings", summary.get("validation", {}).get("finding_count", 0)])
    valuation = summary.get("valuation", {})
    for label, key in [
        ("Baseline expected cost", "baseline_expected_cost"),
        ("Defect count", "defect_count"),
        ("Material defect count", "material_defect_count"),
        ("Combined impact", "combined_impact"),
        ("Sum of isolated impacts", "sum_isolated_impacts"),
        ("Interaction residual", "interaction_residual"),
        ("Materiality threshold (%)", "materiality_threshold_percent"),
    ]:
        ws.append([label, valuation.get(key, "")])
    ws.append(["Interpretation", summary.get("interpretation", "")])
    _style(ws)
    _records_sheet(wb, "Validation Findings", findings)
    _records_sheet(wb, "Valuation Impacts", impacts)
    wb.save(path)
    return path


def _records_sheet(wb, title, records):
    ws = wb.create_sheet(title)
    if not records:
        ws.append(["No records"])
    else:
        headers = list(records[0].keys())
        ws.append(headers)
        for record in records:
            ws.append([record.get(header, "") for header in headers])
    _style(ws)


def _style(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions if ws.max_row > 1 else "A1"
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="243B53")
    for column in ws.columns:
        letter = get_column_letter(column[0].column)
        width = max((len(str(cell.value or "")) for cell in column), default=10)
        ws.column_dimensions[letter].width = min(max(width + 2, 12), 55)


def export_pdf(summary: dict[str, Any], findings: list[dict], impacts: list[dict], path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(path), pagesize=landscape(letter), rightMargin=28, leftMargin=28)
    story = [
        Paragraph("ActuTrace — Analysis Report", styles["Title"]),
        Paragraph("Life-insurance data quality and controlled valuation-impact analysis", styles["Normal"]),
        Spacer(1, 12),
    ]
    meta = summary.get("run_metadata", {})
    val = summary.get("valuation", {})
    rows = [
        ["Run ID", str(meta.get("run_id", "")), "Generated (UTC)", str(meta.get("generated_at_utc", ""))],
        ["Validation findings", str(summary.get("validation", {}).get("finding_count", 0)), "Material defects", str(val.get("material_defect_count", 0))],
        ["Baseline expected cost", str(val.get("baseline_expected_cost", "")), "Combined impact", str(val.get("combined_impact", ""))],
        ["Sum isolated impacts", str(val.get("sum_isolated_impacts", "")), "Interaction residual", str(val.get("interaction_residual", ""))],
    ]
    story += [_table(rows, [110, 190, 110, 250]), Spacer(1, 12)]
    story.append(Paragraph("Validation findings", styles["Heading2"]))
    story.append(_records_table(findings, ["row_number", "policy_id", "field", "rule", "severity", "message"]))
    story.append(Spacer(1, 12))
    story.append(Paragraph("Valuation impacts (ranked)", styles["Heading2"]))
    story.append(_records_table(impacts, ["rank", "defect_id", "policy_id", "field", "original_value", "altered_value", "impact", "absolute_impact", "impact_percent", "is_material"]))
    story.append(Spacer(1, 8))
    story.append(Paragraph("Valuation impacts describe controlled simulated scenarios, not production actuarial estimates. Validation findings do not map one-to-one to financial impacts.", styles["Italic"]))
    doc.build(story)
    return path


def _table(rows, widths):
    table = Table(rows, colWidths=widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#243B53")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D9E2EC")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
    ]))
    return table


def _records_table(records, columns):
    if not records:
        return Paragraph("No records found.", getSampleStyleSheet()["Normal"])
    styles = getSampleStyleSheet()
    rows = [[Paragraph("<b>" + key + "</b>", styles["BodyText"]) for key in columns]]
    for record in records:
        rows.append([Paragraph(str(record.get(key, "") or ""), styles["BodyText"]) for key in columns])
    width = (landscape(letter)[0] - 56) / len(columns)
    return _table(rows, [width] * len(columns))
